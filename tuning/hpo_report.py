"""Adult / Forest-Cover template HPO Excel + notation writer for every dataset.

Workbook structure (3 sheets):
  1. Metric Guide
  2. Performance Metrics  (2-row grouped header, color key)
  3. Best Hyperparameters

Generators included: GaussianCopula, CopulaGAN, CTGAN, TVAE, CTAB-GAN+, WGAN-GP, TabDDPM
(ForestDiffusion excluded until more datasets complete.)
"""

from __future__ import annotations

import json
import re
import warnings
from datetime import date
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sklearn.compose import ColumnTransformer
from sklearn.ensemble import RandomForestClassifier, RandomForestRegressor
from sklearn.linear_model import LinearRegression, LogisticRegression
from sklearn.metrics import accuracy_score, r2_score
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import LabelEncoder, OneHotEncoder, StandardScaler
from sklearn.tree import DecisionTreeClassifier, DecisionTreeRegressor

from config_utils import PACKAGE_ROOT
from datasets.loader import load_dataset_config, load_train_test
from tuning.artifacts import run_dir

warnings.filterwarnings("ignore")

# Core GAN/VAE generators + TabDDPM (all 15 datasets complete).
GENERATORS = [
    ("gaussian_copula", "GaussianCopula"),
    ("copulagan", "CopulaGAN"),
    ("ctgan", "CTGAN"),
    ("tvae", "TVAE"),
    ("ctabgan", "CTAB-GAN+"),
    ("wgan_gp", "WGAN-GP"),
    ("tabddpm", "TabDDPM"),
]
N_GENERATORS = len(GENERATORS)

FILL_DARK = PatternFill("solid", fgColor="1F4E79")
FILL_MAIN = PatternFill("solid", fgColor="2E75B6")
FILL_SUB = PatternFill("solid", fgColor="BDD7EE")
FILL_REF = PatternFill("solid", fgColor="FCE4D6")
FILL_RANK1 = PatternFill("solid", fgColor="C6EFCE")
FILL_MAIN_PALE = PatternFill("solid", fgColor="DDEBF7")
FONT_WHITE = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
FONT_DARK = Font(name="Calibri", bold=True, color="1F4E79", size=10)
FONT_TITLE = Font(name="Calibri", bold=True, size=14, color="1F4E79")
FONT_SECTION = Font(name="Calibri", bold=True, size=12, color="1F4E79")
FONT_BODY = Font(name="Calibri", size=10)
THIN = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

PREFERRED_PARAM_ORDER = [
    "epochs",
    "batch_size",
    "generator_lr",
    "discriminator_lr",
    "generator_dim",
    "discriminator_dim",
    "discriminator_steps",
    "pac",
    "default_distribution",
    "embedding_dim",
    "compress_dims",
    "decompress_dims",
    "l2scale",
    "loss_factor",
    "enforce_min_max_values",
    "latent_dim",
    "hidden_dim",
    "lr",
    "n_critic",
    "lambda_gp",
    "random_dim",
    "num_channels",
    "class_dim",
    "n_t",
    "duplicate_K",
    "diffusion_type",
    "max_depth",
    "n_estimators",
    "eta",
    "num_timesteps",
    "steps",
    "n_layers",
    "d_layers",
    "sample_proportion",
]

SHORT_TITLE = {
    "cancer": "Cancer",
    "alzheimers": "Alzheimers",
    "adult": "Adult Census",
    "forest_cover": "Forest Cover",
    "bank_marketing": "Bank Marketing",
    "wine_quality": "Wine Quality",
    "cdc_diabetes": "CDC Diabetes",
    "mushroom": "Mushroom",
    "magic_gamma": "MAGIC Gamma",
    "metro_interstate": "Metro Interstate",
    "online_shopping": "Online Shopping",
    "air_quality": "Air Quality",
    "concrete": "Concrete",
    "energy_efficiency": "Energy Efficiency",
    "real_estate": "Real Estate",
}

GEN_ALIASES = {
    "gaussiancopula": "gaussian_copula",
    "gaussian_copula": "gaussian_copula",
    "copulagan": "copulagan",
    "ctgan": "ctgan",
    "tvae": "tvae",
    "ctabgan": "ctabgan",
    "ctab-gan+": "ctabgan",
    "ctab-gan": "ctabgan",
    "wgangp": "wgan_gp",
    "wgan-gp": "wgan_gp",
    "wgan_gp": "wgan_gp",
    "forestdiffusion": "forest_diffusion",
    "forest_diffusion": "forest_diffusion",
    "tabddpm": "tabddpm",
}


def _r4(x: Any) -> float | None:
    if x is None or (isinstance(x, float) and (np.isnan(x) or np.isinf(x))):
        return None
    return round(float(x), 4)


def _gen_key(name: str) -> str | None:
    raw = str(name).strip()
    compact = raw.lower().replace(" ", "").replace("_", "")
    if raw.lower() in GEN_ALIASES:
        return GEN_ALIASES[raw.lower()]
    if compact in GEN_ALIASES:
        return GEN_ALIASES[compact]
    for key, display in GENERATORS:
        if raw.lower() == display.lower() or compact == display.lower().replace(" ", "").replace("_", "").replace("+", "+"):
            return key
    return None


def _header_idx(headers: list[str], *needles: str) -> int | None:
    for i, header in enumerate(headers):
        text = header.lower()
        if all(n.lower() in text for n in needles):
            return i
    return None


def _cell_float(row: list[Any], idx: int | None) -> float | None:
    if idx is None or idx >= len(row):
        return None
    val = row[idx]
    if val is None or val == "":
        return None
    try:
        out = float(val)
    except (TypeError, ValueError):
        return None
    if np.isnan(out) or np.isinf(out):
        return None
    return out


def _read_existing_perf(xlsx_path: Path) -> tuple[dict[str, dict[str, float]], tuple[int, int] | None]:
    """Reuse Accuracy/R² columns already written in a previous workbook."""
    if not xlsx_path.is_file():
        return {}, None
    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    except Exception:
        return {}, None
    sizes = None
    if "Metric Guide" in wb.sheetnames:
        guide = wb["Metric Guide"]["A37"].value or ""
        m = re.search(r"Train\s*=\s*(\d+).{0,20}Test\s*=\s*(\d+)", str(guide), flags=re.I)
        if m:
            sizes = (int(m.group(1)), int(m.group(2)))
    if "Performance Metrics" not in wb.sheetnames:
        wb.close()
        return {}, sizes
    ws = wb["Performance Metrics"]
    header_row = None
    if str(ws.cell(1, 1).value or "").strip() == "ID":
        header_row = 2
    else:
        for r in range(1, 8):
            if str(ws.cell(r, 1).value or "").strip() == "Rank":
                header_row = r
                break
    if header_row is None:
        wb.close()
        return {}, sizes
    headers = [
        str(ws.cell(header_row, c).value or "").replace("\n", " ").strip()
        for c in range(1, 31)
    ]
    cols = {
        "generator": _header_idx(headers, "generator"),
        "tstr": _header_idx(headers, "tstr") if _header_idx(headers, "tstr", "acc") is None and _header_idx(headers, "tstr", "r²") is None else (
            _header_idx(headers, "tstr", "acc") or _header_idx(headers, "tstr", "r²") or _header_idx(headers, "tstr")
        ),
        "trtr": _header_idx(headers, "trtr") if _header_idx(headers, "trtr", "acc") is None and _header_idx(headers, "trtr", "r²") is None else (
            _header_idx(headers, "trtr", "acc") or _header_idx(headers, "trtr", "r²") or _header_idx(headers, "trtr")
        ),
        "gap": _header_idx(headers, "gap") if _header_idx(headers, "accuracy gap") is None and _header_idx(headers, "r² gap") is None else (
            _header_idx(headers, "accuracy gap") or _header_idx(headers, "r² gap") or _header_idx(headers, "gap")
        ),
        "tstr_lr": _header_idx(headers, "tstr", "logreg") or _header_idx(headers, "tstr", "linreg") or _header_idx(headers, "tstr lr"),
        "tstr_rf": _header_idx(headers, "tstr", "rf"),
        "tstr_dt": _header_idx(headers, "tstr", "dt"),
        "trtr_lr": _header_idx(headers, "trtr", "logreg") or _header_idx(headers, "trtr", "linreg"),
        "trtr_rf": _header_idx(headers, "trtr", "rf"),
        "trtr_dt": _header_idx(headers, "trtr", "dt"),
        "utility": _header_idx(headers, "utility (accuracy)") or _header_idx(headers, "utility (r²)") or _header_idx(headers, "utility(r²)"),
    }
    # Prefer the standalone TSTR/TRTR mean columns, not per-model ones.
    for name, exacts in (
        ("tstr", ("tstr accuracy", "tstr r²", "tstr r2")),
        ("trtr", ("trtr accuracy", "trtr r²", "trtr r2")),
        ("gap", ("accuracy gap", "r² gap", "r2 gap")),
    ):
        for i, header in enumerate(headers):
            if header.lower() in exacts:
                cols[name] = i
                break
    out: dict[str, dict[str, float]] = {}
    for r in range(header_row + 1, header_row + 20):
        row = [ws.cell(r, c).value for c in range(1, 31)]
        gidx = cols["generator"]
        if gidx is None or not row[gidx]:
            continue
        key = _gen_key(str(row[gidx]))
        if not key:
            continue
        parsed = {name: _cell_float(row, idx) for name, idx in cols.items() if name != "generator"}
        if any(v is not None for v in parsed.values()):
            out[key] = {k: v for k, v in parsed.items() if v is not None}
    wb.close()
    return out, sizes


def _read_accuracy_json(path: Path) -> dict[str, dict[str, float]]:
    if not path.is_file():
        return {}
    raw = json.loads(path.read_text())
    out: dict[str, dict[str, float]] = {}
    clf_map = {
        "logistic_regression": "lr",
        "linear_regression": "lr",
        "random_forest": "rf",
        "decision_tree": "dt",
    }
    for name, block in raw.items():
        key = _gen_key(name)
        if not key or not isinstance(block, dict):
            continue
        row: dict[str, float] = {}
        if "mean_tstr_accuracy" in block:
            row["tstr"] = float(block["mean_tstr_accuracy"])
        if "mean_trtr_accuracy" in block:
            row["trtr"] = float(block["mean_trtr_accuracy"])
        if "accuracy_gap" in block:
            row["gap"] = float(block["accuracy_gap"])
        for src, dest_prefix in (("per_clf_tstr_accuracy", "tstr"), ("per_clf_trtr_accuracy", "trtr")):
            per = block.get(src) or {}
            for clf_name, alias in clf_map.items():
                if clf_name in per:
                    row[f"{dest_prefix}_{alias}"] = float(per[clf_name])
        if row:
            out[key] = row
    return out


def _pipe(X: pd.DataFrame, model):
    X = X.copy()
    for c in X.columns:
        if not pd.api.types.is_numeric_dtype(X[c]):
            X[c] = X[c].astype(str)
    cat = [c for c in X.columns if not pd.api.types.is_numeric_dtype(X[c])]
    num = [c for c in X.columns if c not in cat]
    tf = []
    if num:
        tf.append(("num", StandardScaler(), num))
    if cat:
        tf.append(("cat", OneHotEncoder(handle_unknown="ignore", sparse_output=False), cat))
    return Pipeline([("pre", ColumnTransformer(tf, remainder="drop")), ("est", model)]), X


def _encode_y(*series_list):
    le = LabelEncoder()
    le.fit(pd.concat([s.astype(str) for s in series_list], ignore_index=True))
    return [le.transform(s.astype(str)) for s in series_list]


def _score_split(X_tr, y_tr, X_te, y_te, *, regression: bool) -> dict[str, float]:
    if regression:
        models = {
            "lr": LinearRegression(),
            "rf": RandomForestRegressor(n_estimators=100, max_depth=12, random_state=42, n_jobs=1),
            "dt": DecisionTreeRegressor(max_depth=10, random_state=42),
        }
        metric = r2_score
    else:
        models = {
            "lr": LogisticRegression(max_iter=500, solver="lbfgs"),
            "rf": RandomForestClassifier(n_estimators=100, max_depth=12, random_state=42, n_jobs=1),
            "dt": DecisionTreeClassifier(max_depth=10, random_state=42),
        }
        metric = accuracy_score
    out = {}
    for name, model in models.items():
        pipe, X_tr2 = _pipe(X_tr, model)
        _, X_te2 = _pipe(X_te, model)
        X_te2 = X_te2.reindex(columns=X_tr2.columns)
        pipe.fit(X_tr2, y_tr)
        out[name] = float(metric(y_te, pipe.predict(X_te2)))
    return out


def _score_from_csv(dataset_key: str, synth_path: Path, *, regression: bool) -> dict[str, float] | None:
    train_df, test_df, spec = load_train_test(dataset_key)
    target = spec.require_target()
    synth = pd.read_csv(synth_path)
    for c in train_df.columns:
        if c not in synth.columns:
            synth[c] = np.nan
    synth = synth[train_df.columns].copy()
    for c in train_df.columns:
        if pd.api.types.is_numeric_dtype(train_df[c]):
            synth[c] = pd.to_numeric(synth[c], errors="coerce")
        else:
            synth[c] = synth[c].astype(str)
    X_real = train_df.drop(columns=[target])
    y_real = train_df[target]
    X_test = test_df.drop(columns=[target])
    y_test = test_df[target]
    X_s = synth.drop(columns=[target])
    if regression:
        y_real_e = pd.to_numeric(y_real, errors="coerce").to_numpy()
        y_test_e = pd.to_numeric(y_test, errors="coerce").to_numpy()
        y_s_e = pd.to_numeric(synth[target], errors="coerce").to_numpy()
    else:
        y_real_e, y_test_e, y_s_e = _encode_y(y_real, y_test, synth[target])
    tstr = _score_split(X_s, y_s_e, X_test, y_test_e, regression=regression)
    trtr = _score_split(X_real, y_real_e, X_test, y_test_e, regression=regression)
    tstr_mean = float(np.mean(list(tstr.values())))
    trtr_mean = float(np.mean(list(trtr.values())))
    gap = max(trtr_mean - tstr_mean, 0.0)
    return {
        "tstr": tstr_mean,
        "trtr": trtr_mean,
        "gap": gap,
        "tstr_lr": tstr["lr"],
        "tstr_rf": tstr["rf"],
        "tstr_dt": tstr["dt"],
        "trtr_lr": trtr["lr"],
        "trtr_rf": trtr["rf"],
        "trtr_dt": trtr["dt"],
        "utility": 1.0 - float(np.clip(gap, 0.0, 1.0)),
        "n_train": float(len(train_df)),
        "n_test": float(len(test_df)),
    }


def _collect_rows(dataset_key: str) -> tuple[list[dict], list[str], dict[str, Any]]:
    cfg = load_dataset_config()[dataset_key]
    task = str(cfg.get("task", "classification"))
    regression = task == "regression"
    results_root = PACKAGE_ROOT / "results"
    out_dir = run_dir(results_root, dataset_key, "gaussian_copula", task=task).parent
    existing, sizes = _read_existing_perf(out_dir / f"{dataset_key}_HPO_results.xlsx")
    acc_json = _read_accuracy_json(out_dir / "accuracy_metrics.json")

    rows: list[dict] = []
    all_param_keys: list[str] = []
    n_train, n_test = (sizes or (800, 200))
    for key, display in GENERATORS:
        d = run_dir(results_root, dataset_key, key, task=task)
        if not (d / "completed.json").is_file() or not (d / "metrics.json").is_file():
            continue
        metrics = json.loads((d / "metrics.json").read_text())
        params = json.loads((d / "best_params.json").read_text()) if (d / "best_params.json").is_file() else {}
        timing = json.loads((d / "timing.json").read_text()) if (d / "timing.json").is_file() else {}
        completed = json.loads((d / "completed.json").read_text())

        util_block = existing.get(key) or acc_json.get(key) or {}
        if not util_block and (d / "best_synthetic.csv").is_file():
            try:
                scored = _score_from_csv(dataset_key, d / "best_synthetic.csv", regression=regression)
            except Exception:
                scored = None
            if scored:
                n_train = int(scored.pop("n_train"))
                n_test = int(scored.pop("n_test"))
                util_block = scored

        F = float(metrics["fidelity"])
        P = float(metrics["privacy_risk"])
        tstr_ref = float(metrics.get("mean_tstr_f1", metrics.get("mean_tstr_r2", np.nan)))
        trtr_ref = float(metrics.get("mean_trtr_f1", metrics.get("mean_trtr_r2", np.nan)))
        ref_gap = float(metrics.get("mean_f1_gap", metrics.get("mean_r2_gap", np.nan)))
        if util_block.get("tstr") is not None and util_block.get("trtr") is not None:
            tstr_mean = float(util_block["tstr"])
            trtr_mean = float(util_block["trtr"])
            gap = float(util_block["gap"]) if util_block.get("gap") is not None else max(trtr_mean - tstr_mean, 0.0)
            util = float(util_block["utility"]) if util_block.get("utility") is not None else 1.0 - float(np.clip(max(trtr_mean - tstr_mean, 0.0), 0.0, 1.0))
        else:
            tstr_mean = tstr_ref
            trtr_mean = trtr_ref
            gap = ref_gap if not (isinstance(ref_gap, float) and np.isnan(ref_gap)) else max(trtr_mean - tstr_mean, 0.0)
            util = float(metrics.get("utility", 1.0 - float(np.clip(max(gap, 0.0), 0.0, 1.0))))
        obj = 0.4 * F + 0.4 * util - 0.2 * P
        for k in params:
            if k not in all_param_keys:
                all_param_keys.append(k)
        time_s = timing.get("elapsed_sec")
        if time_s is None:
            time_s = completed.get("elapsed_sec", 0)
        n_trials = int(timing.get("n_trials", 0) or 0)
        rows.append(
            {
                "key": key,
                "generator": display,
                "obj": obj,
                "fidelity": F,
                "utility": util,
                "privacy_risk": P,
                "ks": float(metrics["ks_similarity"]),
                "corr": float(metrics["corr_similarity"]),
                "tstr": tstr_mean,
                "trtr": trtr_mean,
                "gap": gap,
                "tstr_lr": util_block.get("tstr_lr"),
                "tstr_rf": util_block.get("tstr_rf"),
                "tstr_dt": util_block.get("tstr_dt"),
                "trtr_lr": util_block.get("trtr_lr"),
                "trtr_rf": util_block.get("trtr_rf"),
                "trtr_dt": util_block.get("trtr_dt"),
                "mia_auc": float(metrics["mia_auc"]),
                "nndr": float(metrics["nndr"]),
                "dcr": float(metrics.get("dcr", np.nan)),
                "utility_ref": float(metrics["utility"]),
                "tstr_ref": tstr_ref,
                "trtr_ref": trtr_ref,
                "ref_gap": ref_gap,
                "obj_hpo": float(completed.get("best_value", metrics.get("objective", np.nan))),
                "time_s": float(time_s or 0),
                "n_trials": n_trials,
                "params": params,
            }
        )
    rows.sort(key=lambda r: r["obj"], reverse=True)
    param_cols = [k for k in PREFERRED_PARAM_ORDER if k in all_param_keys]
    param_cols += [k for k in all_param_keys if k not in param_cols]
    meta = {
        "dataset_key": dataset_key,
        "title": SHORT_TITLE.get(dataset_key, dataset_key),
        "task": task,
        "regression": regression,
        "n_train": n_train,
        "n_test": n_test,
        "target": str(cfg.get("target", "")),
        "display_name": str(cfg.get("display_name", dataset_key)),
    }
    return rows, param_cols, meta


def _write_metric_guide(ws, meta: dict, n_done: int) -> None:
    regression = meta["regression"]
    util_name = "R²" if regression else "Accuracy"
    title = f"{meta['title']} — Which metrics are in the Objective vs Sub-metrics"
    ws.merge_cells("A1:E1")
    ws["A1"] = title
    ws["A1"].font = FONT_TITLE
    ws.merge_cells("A3:E3")
    ws["A3"] = "OBJECTIVE FORMULA (maximize)"
    ws["A3"].font = FONT_SECTION
    ws.merge_cells("A4:E4")
    ws["A4"] = (
        f"Objective = 0.4 × Fidelity  +  0.4 × Utility ({util_name})  −  0.2 × Privacy Risk"
    )
    ws["A4"].font = Font(name="Calibri", bold=True, size=11)

    role_headers = ["COLOR / ROLE", "Meaning", "In Objective?", "Weight", "Notes"]
    role_rows = [
        ["OBJECTIVE", "Final scalar score used for ranking", "YES — result", "—", "Sort generators by this"],
        ["MAIN — Fidelity", "Statistical similarity", "YES", "+0.4", "Higher better"],
        [f"MAIN — Utility ({util_name})", "ML usefulness via " + util_name, "YES", "+0.4", "Higher better"],
        ["MAIN — Privacy Risk", "Leakage / memorization risk", "YES", "−0.2", "Higher worse (subtracted)"],
        ["SUB → Fidelity", "KS Similarity, Corr Similarity", "NO", "—", "Support / explain Fidelity"],
        [
            "SUB → Utility",
            "TSTR/TRTR R², Gap, per-regressor" if regression else "TSTR/TRTR Acc, Gap, per-classifier",
            "NO",
            "—",
            "Support / explain Utility",
        ],
        ["SUB → Privacy", "MIA AUC, NNDR, DCR", "NO", "—", "Support / explain Privacy Risk"],
        [
            "REFERENCE ONLY",
            "HPO-search metrics (original objective)" if regression else "F1 metrics & original Optuna F1-HPO",
            "NO",
            "—",
            "Not used in current Objective",
        ],
    ]
    for c, h in enumerate(role_headers, 1):
        cell = ws.cell(6, c, h)
        cell.fill = FILL_DARK
        cell.font = FONT_WHITE
        cell.border = THIN
        cell.alignment = CENTER
    for i, row in enumerate(role_rows):
        for c, v in enumerate(row, 1):
            cell = ws.cell(7 + i, c, v)
            cell.border = THIN
            cell.font = FONT_BODY
            cell.alignment = LEFT if c > 1 else CENTER
            if i == 0:
                cell.fill = FILL_DARK
                cell.font = FONT_WHITE
            elif i <= 3:
                cell.fill = FILL_MAIN
                cell.font = FONT_WHITE
            elif i <= 6:
                cell.fill = FILL_SUB
                cell.font = FONT_DARK
            else:
                cell.fill = FILL_REF
                cell.font = FONT_DARK

    ws.merge_cells("A17:E17")
    ws["A17"] = "DETAILED: every metric and whether it enters the Objective"
    ws["A17"].font = FONT_SECTION
    detail_headers = ["Role", "Metric", "Definition", "In Objective?", "How it enters"]
    if regression:
        detail_rows = [
            ["OBJECTIVE", "Objective", "0.4·Fidelity + 0.4·Utility(R²) − 0.2·PrivacyRisk", "YES", "Final ranking score"],
            ["MAIN", "Fidelity", "avg(KS Similarity, Corr Similarity)", "YES (+0.4)", "Main component"],
            ["SUB → Fidelity", "KS Similarity", "1 − KS distance, averaged over columns", "No", "Builds Fidelity"],
            ["SUB → Fidelity", "Corr Similarity", "Correlation matrix similarity real vs synthetic", "No", "Builds Fidelity"],
            ["MAIN", "Utility (R²)", "1 − clip(max(TRTR_R² − TSTR_R², 0), 0, 1)", "YES (+0.4)", "Main component"],
            ["SUB → Utility", "TSTR R²", "Train on Synthetic, Test on Real (mean of 3 models)", "No", "Builds Utility"],
            ["SUB → Utility", "TRTR R²", "Train on Real, Test on Real baseline", "No", "Builds Utility"],
            ["SUB → Utility", "R² Gap", "TRTR R² − TSTR R²", "No", "Builds Utility"],
            ["SUB → Utility", "TSTR/TRTR R² per model", "LinearReg, RF, Decision Tree", "No", "Builds Utility"],
            ["MAIN", "Privacy Risk", "0.5·MIA_AUC + 0.3·NNDR_norm + 0.2·DCR_norm", "YES (−0.2)", "Main component (subtracted)"],
            ["SUB → Privacy", "MIA AUC", "Membership inference attack AUC", "No", "Builds Privacy Risk"],
            ["SUB → Privacy", "NNDR", "Nearest-neighbor distance ratio", "No", "Builds Privacy Risk"],
            ["SUB → Privacy", "DCR", "Distance to closest record", "No", "Builds Privacy Risk"],
            ["REFERENCE", "Utility (HPO) / TSTR / TRTR / Gap", "Metric Optuna used during search", "No", "Reference only"],
            ["REFERENCE", "Objective (HPO)", "Score Optuna maximized during search", "No", "Reference only"],
        ]
    else:
        detail_rows = [
            ["OBJECTIVE", "Objective", "0.4·Fidelity + 0.4·Utility(Accuracy) − 0.2·PrivacyRisk", "YES", "Final ranking score"],
            ["MAIN", "Fidelity", "avg(KS Similarity, Corr Similarity)", "YES (+0.4)", "Main component"],
            ["SUB → Fidelity", "KS Similarity", "1 − KS distance, averaged over columns", "No", "Builds Fidelity"],
            ["SUB → Fidelity", "Corr Similarity", "Correlation matrix similarity real vs synthetic", "No", "Builds Fidelity"],
            ["MAIN", "Utility (Accuracy)", "1 − clip(max(TRTR_Acc − TSTR_Acc, 0), 0, 1)", "YES (+0.4)", "Main component"],
            ["SUB → Utility", "TSTR Accuracy", "Train on Synthetic, Test on Real (mean of 3 clfs)", "No", "Builds Utility"],
            ["SUB → Utility", "TRTR Accuracy", "Train on Real, Test on Real baseline", "No", "Builds Utility"],
            ["SUB → Utility", "Accuracy Gap", "TRTR Accuracy − TSTR Accuracy", "No", "Builds Utility"],
            ["SUB → Utility", "TSTR/TRTR Acc per clf", "LogReg, RF, Decision Tree", "No", "Builds Utility"],
            ["MAIN", "Privacy Risk", "0.5·MIA_AUC + 0.3·NNDR_norm + 0.2·DCR_norm", "YES (−0.2)", "Main component (subtracted)"],
            ["SUB → Privacy", "MIA AUC", "Membership inference attack AUC", "No", "Builds Privacy Risk"],
            ["SUB → Privacy", "NNDR", "Nearest-neighbor distance ratio", "No", "Builds Privacy Risk"],
            ["SUB → Privacy", "DCR", "Distance to closest record", "No", "Builds Privacy Risk"],
            ["REFERENCE", "Utility (F1) / TSTR F1 / TRTR F1 / F1 Gap", "F1-based utility used during Optuna HPO", "No", "Reference only"],
            ["REFERENCE", "Objective (F1 HPO)", "Score Optuna maximized during search", "No", "Reference only"],
        ]
    for c, h in enumerate(detail_headers, 1):
        cell = ws.cell(18, c, h)
        cell.fill = FILL_DARK
        cell.font = FONT_WHITE
        cell.border = THIN
        cell.alignment = CENTER
    for i, row in enumerate(detail_rows):
        for c, v in enumerate(row, 1):
            cell = ws.cell(19 + i, c, v)
            cell.border = THIN
            cell.font = FONT_BODY
            role = row[0]
            if role == "OBJECTIVE":
                cell.fill = FILL_DARK
                cell.font = FONT_WHITE
            elif role == "MAIN":
                cell.fill = FILL_MAIN_PALE
            elif role.startswith("SUB"):
                cell.fill = FILL_MAIN_PALE
            else:
                cell.fill = FILL_REF

    names = ", ".join(d for _, d in GENERATORS)
    ws.merge_cells("A36:E36")
    ws["A36"] = "Scope note"
    ws["A36"].font = FONT_SECTION
    ws.merge_cells("A37:E39")
    pending = N_GENERATORS - n_done
    extra = (
        f" All {N_GENERATORS} generators complete."
        if pending <= 0
        else f" {n_done} of {N_GENERATORS} generators complete; remaining runs still in progress."
    )
    models = "Linear Regression, Random Forest, Decision Tree" if regression else "Logistic Regression, Random Forest, Decision Tree"
    ws["A37"] = (
        f"{meta['title']} HPO report. Generators: {names}.{extra} "
        f"Sampling: max 1000 rows, seed 42. Train={meta['n_train']}, Test={meta['n_test']}. "
        f"Utility models: {models}. Report date: {date.today()}."
    )
    ws["A37"].font = FONT_BODY
    ws["A37"].alignment = LEFT
    for col, w in zip("ABCDE", [22, 28, 55, 16, 28]):
        ws.column_dimensions[col].width = w


def _write_performance(ws, rows: list[dict], meta: dict) -> None:
    regression = meta["regression"]
    util_name = "R²" if regression else "Accuracy"
    top = [
        "ID", None, "OBJECTIVE", "MAIN METRICS (used in Objective)", None, None,
        "SUB → Fidelity", None,
        f"SUB → Utility ({util_name})", None, None, None, None, None, None, None, None,
        "SUB → Privacy Risk", None, None,
        "REFERENCE ONLY (NOT in Objective)", None, None, None, None, "Run",
    ]
    if regression:
        headers = [
            "Rank", "Generator", "Objective\n0.4F+0.4U−0.2P",
            "Fidelity\n(+0.4)", "Utility (R²)\n(+0.4)", "Privacy Risk\n(−0.2)",
            "KS Similarity", "Corr Similarity",
            "TSTR R²", "TRTR R²", "R² Gap",
            "TSTR R² (LinReg)", "TSTR R² (RF)", "TSTR R² (DT)",
            "TRTR R² (LinReg)", "TRTR R² (RF)", "TRTR R² (DT)",
            "MIA AUC", "NNDR", "DCR",
            "Utility (HPO)", "TSTR (HPO)", "TRTR (HPO)", "HPO Gap", "Objective (HPO)",
            "Time (s)",
        ]
    else:
        headers = [
            "Rank", "Generator", "Objective\n0.4F+0.4U−0.2P",
            "Fidelity\n(+0.4)", "Utility (Accuracy)\n(+0.4)", "Privacy Risk\n(−0.2)",
            "KS Similarity", "Corr Similarity",
            "TSTR Accuracy", "TRTR Accuracy", "Accuracy Gap",
            "TSTR Acc (LogReg)", "TSTR Acc (RF)", "TSTR Acc (DT)",
            "TRTR Acc (LogReg)", "TRTR Acc (RF)", "TRTR Acc (DT)",
            "MIA AUC", "NNDR", "DCR",
            "Utility (F1)", "TSTR F1", "TRTR F1", "F1 Gap", "Objective (F1 HPO)",
            "Time (s)",
        ]
    for col, val in enumerate(top, 1):
        cell = ws.cell(1, col, val)
        cell.alignment = CENTER
        cell.border = THIN
    ws.merge_cells("A1:B1")
    ws.merge_cells("D1:F1")
    ws.merge_cells("G1:H1")
    ws.merge_cells("I1:Q1")
    ws.merge_cells("R1:T1")
    ws.merge_cells("U1:Y1")
    for col in range(1, 27):
        cell = ws.cell(1, col)
        if col in (1, 2, 3, 26):
            cell.fill = FILL_DARK
            cell.font = FONT_WHITE
        elif col in (4, 5, 6):
            cell.fill = FILL_MAIN
            cell.font = FONT_WHITE
        elif col in (21, 22, 23, 24, 25):
            cell.fill = FILL_REF
            cell.font = FONT_DARK
        else:
            cell.fill = FILL_SUB
            cell.font = FONT_DARK
    for col, val in enumerate(headers, 1):
        cell = ws.cell(2, col, val)
        cell.fill = FILL_DARK
        cell.font = FONT_WHITE
        cell.alignment = CENTER
        cell.border = THIN

    for i, r in enumerate(rows, start=1):
        vals = [
            i, r["generator"], _r4(r["obj"]),
            _r4(r["fidelity"]), _r4(r["utility"]), _r4(r["privacy_risk"]),
            _r4(r["ks"]), _r4(r["corr"]),
            _r4(r["tstr"]), _r4(r["trtr"]), _r4(r["gap"]),
            _r4(r["tstr_lr"]), _r4(r["tstr_rf"]), _r4(r["tstr_dt"]),
            _r4(r["trtr_lr"]), _r4(r["trtr_rf"]), _r4(r["trtr_dt"]),
            _r4(r["mia_auc"]), _r4(r["nndr"]), _r4(r["dcr"]),
            _r4(r["utility_ref"]), _r4(r["tstr_ref"]), _r4(r["trtr_ref"]), _r4(r["ref_gap"]),
            _r4(r["obj_hpo"]), round(r["time_s"], 1),
        ]
        for col, val in enumerate(vals, 1):
            cell = ws.cell(i + 2, col, val)
            cell.border = THIN
            cell.alignment = CENTER
            cell.font = FONT_BODY
            if i == 1:
                cell.fill = FILL_RANK1

    key_row = 3 + len(rows) + 1
    note_row = key_row + 1
    ws.cell(key_row, 1, "Color key:").font = FONT_DARK
    labels = [
        (2, "OBJECTIVE", FILL_DARK, FONT_WHITE),
        (3, "MAIN (in formula)", FILL_MAIN, FONT_WHITE),
        (4, "SUB-METRIC", FILL_SUB, FONT_DARK),
        (5, "REFERENCE only", FILL_REF, FONT_DARK),
    ]
    for col, text, fill, font in labels:
        cell = ws.cell(key_row, col, text)
        cell.fill = fill
        cell.font = font
        cell.border = THIN
        cell.alignment = CENTER
    formula = (
        "Utility(R²)=1−clip(max(TRTR_R²−TSTR_R²,0),0,1)."
        if regression
        else "Utility(Accuracy)=1−clip(max(TRTR_Acc−TSTR_Acc,0),0,1)."
    )
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=26)
    ws.cell(
        note_row, 1,
        f"{formula} {meta['title']} report with {len(rows)} generators. See Metric Guide.",
    ).font = FONT_BODY

    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 40
    ws.freeze_panes = "C3"
    for col in range(1, 27):
        ws.column_dimensions[get_column_letter(col)].width = 13
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 16


def _write_hyperparams(ws, rows: list[dict], param_cols: list[str]) -> None:
    headers = ["Generator"] + param_cols
    for col, h in enumerate(headers, 1):
        cell = ws.cell(1, col, h)
        cell.fill = FILL_DARK
        cell.font = FONT_WHITE
        cell.alignment = CENTER
        cell.border = THIN
    for i, r in enumerate(rows, start=2):
        ws.cell(i, 1, r["generator"]).border = THIN
        ws.cell(i, 1).font = Font(name="Calibri", bold=True, size=10)
        if i == 2:
            ws.cell(i, 1).fill = FILL_RANK1
        for j, pk in enumerate(param_cols, start=2):
            val = r["params"].get(pk)
            if isinstance(val, (list, tuple)):
                val = str(list(val))
            cell = ws.cell(i, j, val)
            cell.border = THIN
            cell.alignment = CENTER
            if i == 2:
                cell.fill = FILL_RANK1
    ws.column_dimensions["A"].width = 16
    for col in range(2, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16
    ws.freeze_panes = "B2"


def _write_notation(path: Path, rows: list[dict], param_cols: list[str], meta: dict) -> None:
    regression = meta["regression"]
    util_tag = "R²" if regression else "Acc"
    included = ", ".join(r["generator"] for r in rows)
    missing = [d for k, d in GENERATORS if k not in {r["key"] for r in rows}]
    lines = [
        "=" * 80,
        f"      {meta['title'].upper()} — HYPERPARAMETER TUNING RESULTS (Notation File)",
        "=" * 80,
        f"Dataset: {meta['display_name']}",
        f"Task: {'Regression' if regression else 'Classification'} ({meta['target']})",
        "Tuning Framework: Optuna (TPE Sampler, MedianPruner)",
        f"Train/Test sizes after sampling+split: {meta['n_train']} / {meta['n_test']}",
        f"Objective (report ranking): MAXIMIZE  0.4×Fidelity + 0.4×Utility({util_tag}) − 0.2×PrivacyRisk",
        "Note: Optuna search used its own utility; report utility is recomputed as above.",
        "",
        f"SCOPE: {len(rows)} of {N_GENERATORS} generators in this workbook",
        f"  Included: {included}",
    ]
    if missing:
        lines.append(f"  Pending: {', '.join(missing)}")
    lines += [
        "",
        "=" * 80,
        f"                RESULTS RANKED BY {util_tag} OBJECTIVE (Best → Worst)",
        "=" * 80,
    ]
    for i, r in enumerate(rows, 1):
        mins = r["time_s"] / 60.0
        lines += [
            "",
            f"#{i}  {r['generator']:<28} Objective({util_tag}): {r['obj']:.4f}",
            "    " + "─" * 49,
            f"    Fidelity:      {r['fidelity']:.4f}    |  KS Sim: {r['ks']:.4f}   Corr Sim: {r['corr']:.4f}",
            f"    Utility({util_tag}):  {r['utility']:.4f}    |  TSTR: {r['tstr']:.4f}   TRTR: {r['trtr']:.4f}",
            f"    Privacy Risk:  {r['privacy_risk']:.4f}    |  MIA AUC: {r['mia_auc']:.4f}   NNDR: {r['nndr']:.4f}",
            f"    Time: {r['time_s']:.0f}s ({mins:.1f} min)  |  HPO trials: {r['n_trials']}",
            f"    Reference HPO objective: {r['obj_hpo']:.4f}",
            "",
            "    Best Hyperparameters:",
        ]
        for pk in param_cols:
            if pk in r["params"]:
                lines.append(f"      {pk:<20}= {r['params'][pk]}")
    winner = rows[0]
    best_fid = max(rows, key=lambda r: r["fidelity"])
    best_u = max(rows, key=lambda r: r["utility"])
    best_p = min(rows, key=lambda r: r["privacy_risk"])
    lines += [
        "",
        "=" * 80,
        "                        KEY OBSERVATIONS",
        "=" * 80,
        "",
        f"  1. Winner by {util_tag} objective: {winner['generator']} ({winner['obj']:.4f}).",
        f"  2. Top fidelity: {best_fid['generator']} ({best_fid['fidelity']:.4f}).",
        f"  3. Best {util_tag} utility: {best_u['generator']} ({best_u['utility']:.4f}).",
        f"  4. Lowest privacy risk: {best_p['generator']} ({best_p['privacy_risk']:.4f}).",
        "",
        "=" * 80,
        f"  Generated: {date.today()} | Framework: Optuna TPE | Dataset: {meta['dataset_key']} | Generators: {len(rows)}",
        "=" * 80,
        "",
    ]
    path.write_text("\n".join(lines))


def write_dataset(dataset_key: str) -> tuple[Path, Path]:
    rows, param_cols, meta = _collect_rows(dataset_key)
    if not rows:
        raise ValueError(f"No completed generators for {dataset_key}")
    out_dir = run_dir(PACKAGE_ROOT / "results", dataset_key, "gaussian_copula", task=meta["task"]).parent
    xlsx_path = out_dir / f"{dataset_key}_HPO_results.xlsx"
    txt_path = out_dir / f"{dataset_key}_HPO_notation.txt"

    wb = openpyxl.Workbook()
    ws0 = wb.active
    ws0.title = "Metric Guide"
    _write_metric_guide(ws0, meta, n_done=len(rows))
    ws1 = wb.create_sheet("Performance Metrics", 1)
    _write_performance(ws1, rows, meta)
    ws2 = wb.create_sheet("Best Hyperparameters", 2)
    _write_hyperparams(ws2, rows, param_cols)
    wb.save(xlsx_path)
    _write_notation(txt_path, rows, param_cols, meta)
    return xlsx_path, txt_path
