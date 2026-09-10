#!/usr/bin/env python3
"""Write wine_quality HPO notation + Excel (cancer template)."""

from __future__ import annotations

import json
import sys
import warnings
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sklearn.compose import ColumnTransformer
from sklearn.ensemble import RandomForestClassifier
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import accuracy_score
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import LabelEncoder, OneHotEncoder, StandardScaler
from sklearn.tree import DecisionTreeClassifier

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT))
warnings.filterwarnings("ignore")

from datasets.loader import load_train_test  # noqa: E402

DIR = Path(__file__).resolve().parent
GENERATORS = [
    ("gaussian_copula", "GaussianCopula"),
    ("copulagan", "CopulaGAN"),
    ("ctgan", "CTGAN"),
    ("tvae", "TVAE"),
    ("ctabgan", "CTAB-GAN+"),
    ("wgan_gp", "WGAN-GP"),
]

FILL_DARK = PatternFill("solid", fgColor="1F4E79")
FILL_MAIN = PatternFill("solid", fgColor="2E75B6")
FILL_SUB = PatternFill("solid", fgColor="BDD7EE")
FILL_REF = PatternFill("solid", fgColor="FCE4D6")
FILL_RANK1 = PatternFill("solid", fgColor="C6EFCE")
FONT_WHITE = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
FONT_DARK = Font(name="Calibri", bold=True, color="1F4E79", size=10)
FONT_TITLE = Font(name="Calibri", bold=True, size=14, color="1F4E79")
FONT_SECTION = Font(name="Calibri", bold=True, size=12, color="1F4E79")
FONT_BODY = Font(name="Calibri", size=10)
THIN = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

PREFERRED_PARAM_ORDER = [
    "epochs",
    "batch_size",
    "generator_lr",
    "discriminator_lr",
    "generator_dim",
    "discriminator_dim",
    "discriminator_steps",
    "pac",
    "default_distribution",
    "embedding_dim",
    "compress_dims",
    "decompress_dims",
    "l2scale",
    "loss_factor",
    "enforce_min_max_values",
    "latent_dim",
    "hidden_dim",
    "lr",
    "n_critic",
    "lambda_gp",
    "random_dim",
    "num_channels",
    "class_dim",
]


def make_pipe(X, clf):
    X = X.copy()
    for c in X.columns:
        if not pd.api.types.is_numeric_dtype(X[c]):
            X[c] = X[c].astype(str)
    cat = [c for c in X.columns if not pd.api.types.is_numeric_dtype(X[c])]
    num = [c for c in X.columns if c not in cat]
    tf = []
    if num:
        tf.append(("num", StandardScaler(), num))
    if cat:
        tf.append(
            ("cat", OneHotEncoder(handle_unknown="ignore", sparse_output=False), cat)
        )
    return Pipeline([("pre", ColumnTransformer(tf, remainder="drop")), ("clf", clf)]), X


def encode_y(*series_list):
    le = LabelEncoder()
    le.fit(pd.concat([s.astype(str) for s in series_list], ignore_index=True))
    return [le.transform(s.astype(str)) for s in series_list]


CLFS = {
    "lr": LogisticRegression(max_iter=500, solver="lbfgs"),
    "rf": RandomForestClassifier(
        n_estimators=100, max_depth=12, random_state=42, n_jobs=1
    ),
    "dt": DecisionTreeClassifier(max_depth=10, random_state=42),
}


def score_split(X_tr, y_tr, X_te, y_te):
    out = {}
    for name, clf0 in CLFS.items():
        pipe, X_tr2 = make_pipe(X_tr, clf0)
        _, X_te2 = make_pipe(X_te, clf0)
        X_te2 = X_te2.reindex(columns=X_tr2.columns)
        pipe.fit(X_tr2, y_tr)
        out[name] = float(accuracy_score(y_te, pipe.predict(X_te2)))
    return out


def r4(x):
    if x is None or (isinstance(x, float) and (np.isnan(x) or np.isinf(x))):
        return None
    return round(float(x), 4)


def main() -> None:
    missing = [k for k, _ in GENERATORS if not (DIR / k / "completed.json").is_file()]
    if missing:
        raise SystemExit(f"Incomplete generators: {missing}")

    train_df, test_df, spec = load_train_test("wine_quality")
    target = spec.require_target()

    X_real = train_df.drop(columns=[target])
    y_real = train_df[target]
    X_test = test_df.drop(columns=[target])
    y_test = test_df[target]
    y_real_e, y_test_e = encode_y(y_real, y_test)
    trtr = score_split(X_real, y_real_e, X_test, y_test_e)
    trtr_acc_mean = float(np.mean(list(trtr.values())))

    rows = []
    all_param_keys: list[str] = []
    for key, display in GENERATORS:
        d = DIR / key
        metrics = json.loads((d / "metrics.json").read_text())
        params = json.loads((d / "best_params.json").read_text())
        timing = json.loads((d / "timing.json").read_text())
        completed = json.loads((d / "completed.json").read_text())
        synth = pd.read_csv(d / "best_synthetic.csv")
        for c in train_df.columns:
            if c not in synth.columns:
                synth[c] = np.nan
        synth = synth[train_df.columns].copy()
        for c in train_df.columns:
            if pd.api.types.is_numeric_dtype(train_df[c]):
                synth[c] = pd.to_numeric(synth[c], errors="coerce")
            else:
                synth[c] = synth[c].astype(str)

        X_s = synth.drop(columns=[target])
        y_s_e, y_te_e = encode_y(synth[target], y_test)
        tstr = score_split(X_s, y_s_e, X_test, y_te_e)
        tstr_acc_mean = float(np.mean(list(tstr.values())))
        acc_gap = max(trtr_acc_mean - tstr_acc_mean, 0.0)
        util_acc = 1.0 - float(np.clip(acc_gap, 0.0, 1.0))
        F = float(metrics["fidelity"])
        P = float(metrics["privacy_risk"])
        obj_acc = 0.4 * F + 0.4 * util_acc - 0.2 * P
        for k in params:
            if k not in all_param_keys:
                all_param_keys.append(k)
        rows.append(
            {
                "key": key,
                "generator": display,
                "obj_acc": obj_acc,
                "fidelity": F,
                "utility_acc": util_acc,
                "privacy_risk": P,
                "ks": float(metrics["ks_similarity"]),
                "corr": float(metrics["corr_similarity"]),
                "tstr_acc": tstr_acc_mean,
                "trtr_acc": trtr_acc_mean,
                "acc_gap": acc_gap,
                "tstr_acc_lr": tstr["lr"],
                "tstr_acc_rf": tstr["rf"],
                "tstr_acc_dt": tstr["dt"],
                "trtr_acc_lr": trtr["lr"],
                "trtr_acc_rf": trtr["rf"],
                "trtr_acc_dt": trtr["dt"],
                "mia_auc": float(metrics["mia_auc"]),
                "nndr": float(metrics["nndr"]),
                "dcr": float(metrics.get("dcr", np.nan)),
                "utility_f1": float(metrics["utility"]),
                "tstr_f1": float(metrics["mean_tstr_f1"]),
                "trtr_f1": float(metrics["mean_trtr_f1"]),
                "f1_gap": float(metrics["mean_f1_gap"]),
                "obj_f1_hpo": float(
                    completed.get("best_value", metrics.get("objective", np.nan))
                ),
                "time_s": float(timing.get("elapsed_sec", 0)),
                "params": params,
            }
        )

    rows.sort(key=lambda r: r["obj_acc"], reverse=True)
    param_cols = [k for k in PREFERRED_PARAM_ORDER if k in all_param_keys]
    param_cols += [k for k in all_param_keys if k not in param_cols]

    wb = openpyxl.Workbook()
    ws0 = wb.active
    ws0.title = "Metric Guide"
    ws0.merge_cells("A1:E1")
    ws0["A1"] = "Wine Quality — Which metrics are in the Objective vs Sub-metrics"
    ws0["A1"].font = FONT_TITLE
    ws0.merge_cells("A3:E3")
    ws0["A3"] = "OBJECTIVE FORMULA (maximize)"
    ws0["A3"].font = FONT_SECTION
    ws0.merge_cells("A4:E4")
    ws0["A4"] = (
        "Objective = 0.4 × Fidelity  +  0.4 × Utility (Accuracy)  −  0.2 × Privacy Risk"
    )
    ws0["A4"].font = Font(name="Calibri", bold=True, size=11)

    role_headers = ["COLOR / ROLE", "Meaning", "In Objective?", "Weight", "Notes"]
    role_rows = [
        ["OBJECTIVE", "Final scalar score used for ranking", "YES — result", "—", "Sort generators by this"],
        ["MAIN — Fidelity", "Statistical similarity", "YES", "+0.4", "Higher better"],
        ["MAIN — Utility (Accuracy)", "ML usefulness via Accuracy", "YES", "+0.4", "Higher better"],
        ["MAIN — Privacy Risk", "Leakage / memorization risk", "YES", "−0.2", "Higher worse (subtracted)"],
        ["SUB → Fidelity", "KS Similarity, Corr Similarity", "NO", "—", "Support / explain Fidelity"],
        ["SUB → Utility", "TSTR/TRTR Acc, Gap, per-classifier", "NO", "—", "Support / explain Utility"],
        ["SUB → Privacy", "MIA AUC, NNDR, DCR", "NO", "—", "Support / explain Privacy Risk"],
        ["REFERENCE ONLY", "F1 metrics & original Optuna F1-HPO", "NO", "—", "Not used in current Objective"],
    ]
    for c, h in enumerate(role_headers, 1):
        cell = ws0.cell(6, c, h)
        cell.fill = FILL_DARK
        cell.font = FONT_WHITE
        cell.border = THIN
    for i, row in enumerate(role_rows):
        for c, v in enumerate(row, 1):
            cell = ws0.cell(7 + i, c, v)
            cell.border = THIN
            cell.font = FONT_BODY
            if i == 0:
                cell.fill = FILL_DARK
                cell.font = FONT_WHITE
            elif i <= 3:
                cell.fill = FILL_MAIN
                cell.font = FONT_WHITE
            elif i <= 6:
                cell.fill = FILL_SUB
                cell.font = FONT_DARK
            else:
                cell.fill = FILL_REF
                cell.font = FONT_DARK

    ws0.merge_cells("A17:E17")
    ws0["A17"] = "DETAILED: every metric and whether it enters the Objective"
    ws0["A17"].font = FONT_SECTION
    detail_headers = ["Role", "Metric", "Definition", "In Objective?", "How it enters"]
    detail_rows = [
        ["OBJECTIVE", "Objective", "0.4·Fidelity + 0.4·Utility(Accuracy) − 0.2·PrivacyRisk", "YES", "Final ranking score"],
        ["MAIN", "Fidelity", "avg(KS Similarity, Corr Similarity)", "YES (+0.4)", "Main component"],
        ["SUB → Fidelity", "KS Similarity", "1 − KS distance, averaged over columns", "No", "Builds Fidelity"],
        ["SUB → Fidelity", "Corr Similarity", "Correlation matrix similarity real vs synthetic", "No", "Builds Fidelity"],
        ["MAIN", "Utility (Accuracy)", "1 − clip(max(TRTR_Acc − TSTR_Acc, 0), 0, 1)", "YES (+0.4)", "Main component"],
        ["SUB → Utility", "TSTR Accuracy", "Train on Synthetic, Test on Real (mean of 3 clfs)", "No", "Builds Utility"],
        ["SUB → Utility", "TRTR Accuracy", "Train on Real, Test on Real baseline", "No", "Builds Utility"],
        ["SUB → Utility", "Accuracy Gap", "TRTR Accuracy − TSTR Accuracy", "No", "Builds Utility"],
        ["SUB → Utility", "TSTR Acc (LogReg/RF/DT)", "Per-classifier TSTR accuracy", "No", "Detail under Utility"],
        ["SUB → Utility", "TRTR Acc (LogReg/RF/DT)", "Per-classifier TRTR accuracy", "No", "Detail under Utility"],
        ["MAIN", "Privacy Risk", "Weighted mix of MIA / NNDR / DCR", "YES (−0.2)", "Main component (penalized)"],
        ["SUB → Privacy", "MIA AUC", "Membership Inference Attack AUC (w=0.50)", "No", "Builds Privacy Risk"],
        ["SUB → Privacy", "NNDR", "Nearest Neighbor Distance Ratio (w=0.30)", "No", "Builds Privacy Risk"],
        ["SUB → Privacy", "DCR", "Distance to Closest Record (w=0.20)", "No", "Builds Privacy Risk"],
        ["REFERENCE", "Utility (F1)", "Same gap formula using F1 instead of Accuracy", "No", "Not in Objective"],
        ["REFERENCE", "TSTR/TRTR F1, F1 Gap", "F1-based utility components", "No", "Reference only"],
        ["REFERENCE", "Objective (F1 HPO)", "Original Optuna score (F1 utility during search)", "No", "Historical HPO only"],
    ]
    for c, h in enumerate(detail_headers, 1):
        cell = ws0.cell(18, c, h)
        cell.fill = FILL_DARK
        cell.font = FONT_WHITE
        cell.border = THIN
    for i, row in enumerate(detail_rows):
        for c, v in enumerate(row, 1):
            cell = ws0.cell(19 + i, c, v)
            cell.border = THIN
            cell.font = FONT_BODY
            role = row[0]
            if role == "OBJECTIVE":
                cell.fill = FILL_DARK
                cell.font = FONT_WHITE
            elif role == "MAIN":
                cell.fill = PatternFill("solid", fgColor="DDEBF7")
            elif role.startswith("SUB"):
                cell.fill = PatternFill("solid", fgColor="DEEBF7")
            else:
                cell.fill = FILL_REF

    ws0.merge_cells("A37:E37")
    ws0["A37"] = "Classifiers for Utility (Accuracy): Logistic Regression, Random Forest, Decision Tree"
    ws0["A37"].font = FONT_BODY
    ws0.merge_cells("A38:E38")
    ws0["A38"] = (
        f"Scope: COMPLETE with 6 generators (GaussianCopula, CopulaGAN, CTGAN, TVAE, CTAB-GAN+, WGAN-GP). "
        f"ForestDiffusion/TabDDPM excluded. Sampling: max 1000 rows, balanced classes, seed 42. "
        f"Train={len(train_df)}, Test={len(test_df)}. Preprocess: impute+encode. Report date: {date.today()}."
    )
    ws0["A38"].font = FONT_BODY
    ws0["A38"].alignment = LEFT
    for col, w in zip("ABCDE", [22, 30, 58, 16, 26]):
        ws0.column_dimensions[col].width = w

    ws1 = wb.create_sheet("Performance Metrics", 1)
    top = [
        "ID", None, "OBJECTIVE", "MAIN METRICS (used in Objective)", None, None,
        "SUB → Fidelity", None,
        "SUB → Utility (Accuracy)", None, None, None, None, None, None, None, None,
        "SUB → Privacy Risk", None, None,
        "REFERENCE ONLY (NOT in Objective)", None, None, None, None, "Run",
    ]
    headers = [
        "Rank", "Generator", "Objective\n0.4F+0.4U−0.2P",
        "Fidelity\n(+0.4)", "Utility (Accuracy)\n(+0.4)", "Privacy Risk\n(−0.2)",
        "KS Similarity", "Corr Similarity",
        "TSTR Accuracy", "TRTR Accuracy", "Accuracy Gap",
        "TSTR Acc (LogReg)", "TSTR Acc (RF)", "TSTR Acc (DT)",
        "TRTR Acc (LogReg)", "TRTR Acc (RF)", "TRTR Acc (DT)",
        "MIA AUC", "NNDR", "DCR",
        "Utility (F1)", "TSTR F1", "TRTR F1", "F1 Gap", "Objective (F1 HPO)",
        "Time (s)",
    ]
    for col, val in enumerate(top, 1):
        cell = ws1.cell(1, col, val)
        cell.alignment = CENTER
        cell.border = THIN
    ws1.merge_cells("A1:B1")
    ws1.merge_cells("D1:F1")
    ws1.merge_cells("G1:H1")
    ws1.merge_cells("I1:Q1")
    ws1.merge_cells("R1:T1")
    ws1.merge_cells("U1:Y1")
    for col in range(1, 27):
        cell = ws1.cell(1, col)
        if col in (1, 2, 3, 26):
            cell.fill = FILL_DARK
            cell.font = FONT_WHITE
        elif col in (4, 5, 6):
            cell.fill = FILL_MAIN
            cell.font = FONT_WHITE
        elif col in (21, 22, 23, 24, 25):
            cell.fill = FILL_REF
            cell.font = FONT_DARK
        else:
            cell.fill = FILL_SUB
            cell.font = FONT_DARK
    for col, val in enumerate(headers, 1):
        cell = ws1.cell(2, col, val)
        cell.fill = FILL_DARK
        cell.font = FONT_WHITE
        cell.alignment = CENTER
        cell.border = THIN

    for i, r in enumerate(rows, start=1):
        vals = [
            i, r["generator"], r4(r["obj_acc"]),
            r4(r["fidelity"]), r4(r["utility_acc"]), r4(r["privacy_risk"]),
            r4(r["ks"]), r4(r["corr"]),
            r4(r["tstr_acc"]), r4(r["trtr_acc"]), r4(r["acc_gap"]),
            r4(r["tstr_acc_lr"]), r4(r["tstr_acc_rf"]), r4(r["tstr_acc_dt"]),
            r4(r["trtr_acc_lr"]), r4(r["trtr_acc_rf"]), r4(r["trtr_acc_dt"]),
            r4(r["mia_auc"]), r4(r["nndr"]), r4(r["dcr"]),
            r4(r["utility_f1"]), r4(r["tstr_f1"]), r4(r["trtr_f1"]), r4(r["f1_gap"]),
            r4(r["obj_f1_hpo"]), round(r["time_s"], 1),
        ]
        for col, val in enumerate(vals, 1):
            cell = ws1.cell(i + 2, col, val)
            cell.border = THIN
            cell.alignment = CENTER
            cell.font = FONT_BODY
            if i == 1:
                cell.fill = FILL_RANK1
    ws1.row_dimensions[1].height = 24
    ws1.row_dimensions[2].height = 40
    ws1.freeze_panes = "C3"
    for col in range(1, 27):
        ws1.column_dimensions[get_column_letter(col)].width = 13
    ws1.column_dimensions["B"].width = 16

    ws2 = wb.create_sheet("Best Hyperparameters", 2)
    hp_headers = ["Generator"] + param_cols
    for col, h in enumerate(hp_headers, 1):
        cell = ws2.cell(1, col, h)
        cell.fill = FILL_DARK
        cell.font = FONT_WHITE
        cell.alignment = CENTER
        cell.border = THIN
    for i, r in enumerate(rows, start=2):
        ws2.cell(i, 1, r["generator"]).border = THIN
        ws2.cell(i, 1).font = Font(name="Calibri", bold=True, size=10)
        if i == 2:
            ws2.cell(i, 1).fill = FILL_RANK1
        for j, pk in enumerate(param_cols, start=2):
            val = r["params"].get(pk)
            if isinstance(val, (list, tuple)):
                val = str(list(val))
            cell = ws2.cell(i, j, val)
            cell.border = THIN
            cell.alignment = CENTER
            if i == 2:
                cell.fill = FILL_RANK1
    ws2.column_dimensions["A"].width = 16
    for col in range(2, len(hp_headers) + 1):
        ws2.column_dimensions[get_column_letter(col)].width = 16
    ws2.freeze_panes = "B2"

    xlsx_path = DIR / "wine_quality_HPO_results.xlsx"
    wb.save(xlsx_path)

    def fmt_params(params):
        return "\n".join(f"      {k:<20}= {v}" for k, v in params.items())

    winner = rows[0]
    best_fid = max(rows, key=lambda r: r["fidelity"])
    best_u = max(rows, key=lambda r: r["utility_acc"])
    best_p = min(rows, key=lambda r: r["privacy_risk"])

    notation = f"""================================================================================
      WINE QUALITY — HYPERPARAMETER TUNING RESULTS (Notation File)
================================================================================
Dataset: Wine Quality (white)
Task: Multi-class Classification (quality score)
Tuning Framework: Optuna (TPE Sampler, MedianPruner)
Trials per Generator: 20
Sampling: max_samples=1000, balance_classes=true, seed=42
Train/Test sizes after sampling+split: {len(train_df)} / {len(test_df)}
Preprocessing: impute (median/mode) + categorical label-encoding (train-fitted)
Objective (report ranking): MAXIMIZE  0.4×Fidelity + 0.4×Utility(Accuracy) − 0.2×PrivacyRisk
Note: Optuna search itself used F1-based utility; Acc utility recomputed for this report.

SCOPE: COMPLETE with 6 generators
  Included: GaussianCopula, CopulaGAN, CTGAN, TVAE, CTAB-GAN+, WGAN-GP
  Excluded: ForestDiffusion, TabDDPM

================================================================================
                RESULTS RANKED BY ACCURACY OBJECTIVE (Best → Worst)
================================================================================
"""
    for i, r in enumerate(rows, 1):
        mins = r["time_s"] / 60.0
        notation += f"""
#{i}  {r['generator']:<28} Objective(Acc): {r['obj_acc']:.4f}
    ─────────────────────────────────────────────────
    Fidelity:      {r['fidelity']:.4f}    |  KS Sim: {r['ks']:.4f}   Corr Sim: {r['corr']:.4f}
    Utility(Acc):  {r['utility_acc']:.4f}    |  TSTR Acc: {r['tstr_acc']:.4f}   TRTR Acc: {r['trtr_acc']:.4f}
    Privacy Risk:  {r['privacy_risk']:.4f}    |  MIA AUC: {r['mia_auc']:.4f}   NNDR: {r['nndr']:.4f}
    Time: {r['time_s']:.0f}s ({mins:.1f} min)
    Reference F1-HPO objective: {r['obj_f1_hpo']:.4f}  |  Utility(F1): {r['utility_f1']:.4f}

    Best Hyperparameters:
{fmt_params(r['params'])}
"""
    notation += f"""
================================================================================
                        SKIPPED / EXCLUDED
================================================================================

  ForestDiffusion: Not included in this 6-generator wine_quality sweep.
  TabDDPM:         Dependencies not installed / excluded from phase-1 sweeps.

================================================================================
                        KEY OBSERVATIONS
================================================================================

  1. Winner by Acc objective: {winner['generator']} ({winner['obj_acc']:.4f}).
  2. Top fidelity: {best_fid['generator']} ({best_fid['fidelity']:.4f}).
  3. Best Acc utility: {best_u['generator']} ({best_u['utility_acc']:.4f}).
  4. Lowest privacy risk: {best_p['generator']} ({best_p['privacy_risk']:.4f}).
  5. Ranking uses Accuracy utility; F1-HPO objectives kept as reference only.

================================================================================
  Generated: {date.today()} | Framework: Optuna TPE | Dataset: wine_quality | Generators: 6
================================================================================
"""
    txt_path = DIR / "wine_quality_HPO_notation.txt"
    txt_path.write_text(notation)
    print(f"Wrote {xlsx_path}")
    print(f"Wrote {txt_path}")
    for i, r in enumerate(rows, 1):
        print(
            f"  #{i} {r['generator']:14s} AccObj={r['obj_acc']:.4f} "
            f"F={r['fidelity']:.4f} Uacc={r['utility_acc']:.4f} P={r['privacy_risk']:.4f}"
        )


if __name__ == "__main__":
    main()
